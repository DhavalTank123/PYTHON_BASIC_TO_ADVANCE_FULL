# Excel Read & Write

import openpyxl 
import os

file_name = "student_records.xlsx"

if os.path.exists(file_name):
	workbook = openpyxl.load_workbook(file_name)
	sheet = workbook.active

else:
	workbook = openpyxl.Workbook()
	sheet = workbook.active

	sheet["A1"] = "Roll No"
	sheet["B1"] = "Name"
	sheet["C1"] = "Age"
	sheet["D1"] = "Course"
	sheet["E1"] = "City"

roll_no = int(input("Enter Roll No: "))
name = input("Enter Student Name: ")
age = int(input("Enter Student Age: "))
course = input("Enter Course Name: ")
city = input("Enter City: ")

sheet.append([
	roll_no,
	name,
	age,
	course,
	city
])

workbook.save(file_name)

print("\nStudent record saved successfully!")

print("\n---------- Student Records ------------")

workbook = openpyxl.load_workbook(file_name)
sheet = workbook.active

for row in sheet.iter_rows(values_only=True):
	print(row)